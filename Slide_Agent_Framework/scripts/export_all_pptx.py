import asyncio, os
from playwright.async_api import async_playwright
import pptx

async def export_deck(html_filename, output_pptx_filename):
    html_path = os.path.abspath(html_filename).replace('\\', '/')
    file_url = f'file:///{html_path}'
    out_dir = os.path.abspath('exports')
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, output_pptx_filename)

    print(f'\n--- Exporting {html_filename} to {output_pptx_filename} ---')

    async with async_playwright() as p:
        browser = await p.chromium.launch(
            executable_path=r'C:\Program Files\Google\Chrome\Application\chrome.exe',
            args=['--allow-file-access-from-files']
        )
        page = await browser.new_page(viewport={"width": 1920, "height": 1080})
        
        await page.goto(file_url, wait_until='networkidle')
        await asyncio.sleep(1)

        # Trigger HtmlToPptxConverter
        res = await page.evaluate('''async () => {
            const activeDeck = document.getElementById('slide-stage');
            const slides = Array.from(activeDeck.querySelectorAll('.slide'));
            
            const pptx = new window.PptxGenJS();
            pptx.layout = 'LAYOUT_16x9';

            for (let i = 0; i < slides.length; i++) {
                await window.HtmlToPptxConverter.convertSlideToPptx(slides[i], pptx, i + 1, slides.length);
            }

            const base64 = await pptx.write('base64');
            return {
                slideCount: slides.length,
                base64: base64
            };
        }''')

        import base64
        pptx_bytes = base64.b64decode(res['base64'])
        with open(out_path, 'wb') as f:
            f.write(pptx_bytes)

        print(f'  Saved {out_path} ({len(pptx_bytes)/1024:.1f} KB, {res["slideCount"]} slides)')
        await browser.close()

    # Validate with python-pptx
    prs = pptx.Presentation(out_path)
    print(f'  [VALID PPTX] Loaded {len(prs.slides)} slides successfully.')
    for idx, s in enumerate(prs.slides):
        if idx >= 3: break
        print(f'    Slide {idx+1}: {len(s.shapes)} shapes')

async def main():
    await export_deck('slide_A3_hoan_thien.html', 'On_Tap_Bai_A3_DaiNam_2K.pptx')
    await export_deck('slide_A4_hoan_thien.html', 'On_Tap_Bai_A4_DaiNam_2K.pptx')
    print('\nALL PPTX DECKS EXPORTED AND VALIDATED SUCCESSFULLY!')

if __name__ == '__main__':
    asyncio.run(main())

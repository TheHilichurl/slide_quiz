import sys
from pptx import Presentation
sys.stdout.reconfigure(encoding='utf-8')

pptx_file = sys.argv[1] if len(sys.argv) > 1 else 'exports/slide_A4_hoan_thien_Editable.pptx'
prs = Presentation(pptx_file)
print(f"Presentation: {pptx_file}, slides={len(prs.slides)}")

slide_idx = int(sys.argv[2]) if len(sys.argv) > 2 else 0
slide = prs.slides[slide_idx]
print(f'=== SHAPES & IMAGES ON SLIDE {slide_idx} ===')
for i, s in enumerate(slide.shapes):
    print(f'Shape {i}: name="{s.name}", type={s.shape_type}, left={s.left.inches:.2f}", top={s.top.inches:.2f}", w={s.width.inches:.2f}", h={s.height.inches:.2f}"')
    if s.has_text_frame:
        for p in s.text_frame.paragraphs:
            for r in p.runs:
                print(f'   -> "{r.text[:35]}" ({r.font.name}, {r.font.size.pt if r.font.size else None}pt)')
